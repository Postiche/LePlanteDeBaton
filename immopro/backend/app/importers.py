"""Extraction structurée depuis Excel, CSV, PDF et Word.

Les valeurs extraites sont TOUJOURS soumises à une validation humaine avant enregistrement.
Une valeur non trouvée reste vide : rien n'est déduit ni inventé. Chaque valeur extraite d'un
texte libre est accompagnée de l'extrait source (« preuve ») pour faciliter la vérification.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime

from .geo import normalize

# ---------------------------------------------------------------------------
# Correspondance des colonnes
# ---------------------------------------------------------------------------

PROPERTY_COLUMNS: dict[str, list[str]] = {
    "reference": ["reference", "ref", "ref annonce", "numero", "n mandat", "mandat"],
    "title": ["titre", "intitule", "libelle", "nom du bien"],
    "property_type": ["type", "type de bien", "nature", "typologie", "categorie"],
    "transaction": ["transaction", "vente location", "type de transaction", "operation"],
    "address": ["adresse", "rue", "adresse du bien"],
    "postal_code": ["code postal", "cp", "code_postal"],
    "city": ["ville", "commune", "localite"],
    "latitude": ["latitude", "lat"],
    "longitude": ["longitude", "lon", "lng"],
    "surface_total": ["surface", "surface totale", "surface m2", "surface m²", "surface totale m2", "superficie", "sup totale"],
    "surface_divisible_min": ["divisible", "surface divisible", "divisible a partir de", "surface min divisible"],
    "surface_offices": ["bureaux", "surface bureaux", "bureaux m2"],
    "surface_storage": ["stockage", "surface stockage", "entrepot m2", "surface entrepot", "activite m2"],
    "land_area": ["terrain", "surface terrain", "terrain m2"],
    "clear_height": ["hauteur", "hauteur libre", "hsp", "hauteur sous poutre"],
    "dock_count": ["quais", "nombre de quais", "nb quais", "quai"],
    "sectional_doors": ["portes sectionnelles", "portes", "nb portes", "portes plain pied"],
    "hgv_access": ["acces pl", "acces poids lourds", "pl", "poids lourds"],
    "floor_load": ["resistance au sol", "charge au sol", "resistance sol", "t m2"],
    "power_kva": ["puissance", "puissance electrique", "kva", "puissance kva"],
    "parking_spaces": ["parking", "places de parking", "stationnement", "nb parkings"],
    "condition": ["etat", "etat du bien"],
    "available_from": ["disponibilite", "date de disponibilite", "disponible le", "date dispo"],
    "sale_price": ["prix", "prix de vente", "prix vente", "prix hd", "prix net vendeur"],
    "annual_rent": ["loyer", "loyer annuel", "loyer an", "loyer annuel ht hc", "loyer ht hc an"],
    "rent_per_sqm": ["loyer m2", "loyer m²", "loyer au m2", "loyer au m²", "loyer m2 an", "eur m2 an"],
    "charges": ["charges", "charges annuelles"],
    "property_tax": ["taxe fonciere", "tf"],
    "fees": ["honoraires", "frais d agence"],
    "description": ["description", "descriptif", "commentaire", "commentaires", "observations"],
    "owner_name": ["proprietaire", "bailleur", "vendeur"],
    "agency_name": ["agence", "cabinet", "annonceur"],
    "sales_name": ["commercial", "negociateur", "consultant"],
    "contact_email": ["email", "e mail", "mail", "courriel", "email contact"],
    "contact_phone": ["telephone", "tel", "tel contact", "portable"],
    "source_url": ["url", "lien", "url annonce", "lien annonce"],
}

REQUEST_COLUMNS: dict[str, list[str]] = {
    "reference": ["reference", "ref", "numero"],
    "title": ["titre", "intitule", "objet"],
    "client_name": ["client", "societe", "entreprise", "raison sociale"],
    "contact_name": ["interlocuteur", "contact", "nom du contact", "nom"],
    "email": ["email", "e mail", "mail", "courriel"],
    "phone": ["telephone", "tel", "portable", "mobile"],
    "property_types": ["type", "type de bien", "types de bien", "nature"],
    "transaction": ["transaction", "vente location", "operation"],
    "cities": ["communes", "villes", "ville", "commune", "secteur", "secteurs"],
    "postal_codes": ["codes postaux", "code postal", "cp"],
    "departments": ["departements", "departement", "dept", "dpt"],
    "regions": ["regions", "region"],
    "radius_km": ["rayon", "rayon km", "rayon de recherche"],
    "surface_min": ["surface min", "surface minimale", "surface minimum", "surface mini"],
    "surface_max": ["surface max", "surface maximale", "surface maximum", "surface maxi"],
    "budget_max": ["budget", "budget max", "loyer max", "prix max", "budget maximum"],
    "budget_per_sqm_max": ["budget m2", "loyer m2 max", "prix m2 max"],
    "min_clear_height": ["hauteur", "hauteur libre", "hauteur min"],
    "min_docks": ["quais", "nombre de quais", "quais min"],
    "requires_hgv_access": ["acces pl", "acces poids lourds", "poids lourds"],
    "min_floor_load": ["resistance au sol", "charge au sol"],
    "min_power_kva": ["puissance", "puissance kva"],
    "min_parking": ["parking", "places de parking"],
    "desired_date": ["date souhaitee", "disponibilite souhaitee", "echeance", "date d entree"],
    "priority": ["priorite"],
    "comments": ["commentaires", "commentaire", "observations", "contraintes"],
    "imperatives": ["criteres imperatifs", "imperatifs"],
}

NUMERIC_FIELDS = {
    "latitude", "longitude", "surface_total", "surface_divisible_min", "surface_offices", "surface_storage", "land_area",
    "clear_height", "floor_load", "power_kva", "sale_price", "annual_rent", "rent_per_sqm", "charges", "property_tax",
    "radius_km", "surface_min", "surface_max", "budget_max", "budget_per_sqm_max", "min_clear_height", "min_floor_load", "min_power_kva",
}
INT_FIELDS = {"dock_count", "sectional_doors", "parking_spaces", "min_docks", "min_parking"}
BOOL_FIELDS = {"hgv_access", "requires_hgv_access"}
DATE_FIELDS = {"available_from", "desired_date"}
LIST_FIELDS = {"cities", "postal_codes", "departments", "regions", "property_types", "imperatives"}

TYPE_KEYWORDS = [
    ("plateforme_logistique", ["plateforme logistique", "plate forme logistique", "logistique", "messagerie"]),
    ("entrepot", ["entrepot", "stockage", "depot"]),
    ("local_activite", ["local d activite", "locaux d activite", "local activite", "activite", "atelier"]),
    ("batiment_industriel", ["batiment industriel", "industriel", "usine", "site industriel"]),
    ("bureaux", ["bureaux", "bureau"]),
    ("terrain", ["terrain"]),
    ("commerce", ["commerce", "boutique", "local commercial"]),
]

IMPERATIVE_WORDS = {
    "localisation": ["localisation", "secteur", "zone", "lieu"],
    "surface": ["surface"],
    "budget": ["budget", "prix", "loyer"],
    "typologie": ["type", "typologie"],
    "transaction": ["transaction", "vente", "location"],
    "hauteur_libre": ["hauteur"],
    "quais": ["quai"],
    "acces_pl": ["poids lourds", "pl"],
    "resistance_sol": ["sol", "charge"],
    "puissance": ["puissance", "electrique"],
    "parking": ["parking"],
    "disponibilite": ["disponibilite", "date", "delai"],
}


def parse_number(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().lower()
    s = s.replace(" ", " ").replace("\xa0", " ")
    m = re.search(r"-?\d[\d ]*(?:[.,]\d+)?", s)
    if not m:
        return None
    num = m.group(0).replace(" ", "")
    if "," in num:
        num = num.replace(".", "").replace(",", ".")
    elif num.count(".") > 1 or re.fullmatch(r"\d{1,3}(\.\d{3})+", num):
        num = num.replace(".", "")
    try:
        val = float(num)
    except ValueError:
        return None
    if re.search(r"\d\s*(k€|k ?eur)", s):
        val *= 1000
    elif re.search(r"\d\s*(m€|millions?)", s):
        val *= 1_000_000
    return val


def parse_bool(v) -> bool | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    s = normalize(str(v))
    if s in ("oui", "o", "yes", "y", "x", "vrai", "true", "1", "ok"):
        return True
    if s in ("non", "n", "no", "faux", "false", "0"):
        return False
    return None


def parse_date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    m = re.search(r"(\d{1,2})/(\d{4})", s)
    if m:
        return date(int(m.group(2)), int(m.group(1)), 1).isoformat()
    return None


def parse_types(v) -> list[str]:
    n = normalize(str(v or ""))
    found = []
    for key, words in TYPE_KEYWORDS:
        if any(re.search(rf"\b{w}\b", n) for w in words) and key not in found:
            found.append(key)
    # « entrepôt logistique » → privilégier la plateforme mais garder l'entrepôt
    return found


def parse_transaction(v) -> str | None:
    n = normalize(str(v or ""))
    sale = bool(re.search(r"\b(vente|achat|cession|acquisition|acquerir|a vendre)\b", n))
    rent = bool(re.search(r"\b(location|bail|louer|a louer|locatif|loyer)\b", n))
    if sale and rent:
        return "vente_ou_location"
    return "vente" if sale else "location" if rent else None


def parse_list(v) -> list[str]:
    if v is None or v == "":
        return []
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    return [x.strip() for x in re.split(r"[;,/|\n]+", str(v)) if x.strip()]


def parse_imperatives(v) -> list[str]:
    n = normalize(str(v or ""))
    return [k for k, words in IMPERATIVE_WORDS.items() if any(re.search(rf"\b{w}", n) for w in words)]


def convert_value(field: str, v):
    if field in NUMERIC_FIELDS:
        return parse_number(v)
    if field in INT_FIELDS:
        n = parse_number(v)
        return int(n) if n is not None else None
    if field in BOOL_FIELDS:
        return parse_bool(v)
    if field in DATE_FIELDS:
        return parse_date(v)
    if field == "property_type":
        t = parse_types(v)
        return t[0] if t else None
    if field == "property_types":
        return parse_types(v)
    if field == "transaction":
        return parse_transaction(v)
    if field == "imperatives":
        return parse_imperatives(v)
    if field == "departments":
        return [d.upper().zfill(2) if d.isdigit() else d.upper() for d in parse_list(v)]
    if field in LIST_FIELDS:
        return parse_list(v)
    if field == "postal_code":
        s = str(v).strip()
        if re.fullmatch(r"\d{4}(\.0)?", s):  # Excel supprime le zéro initial
            s = "0" + s[:4]
        s = s.replace(".0", "")
        return s if re.fullmatch(r"\d{5}", s) else None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def map_headers(headers: list, target: str) -> dict[int, str]:
    columns = PROPERTY_COLUMNS if target == "property" else REQUEST_COLUMNS
    mapping: dict[int, str] = {}
    used: set[str] = set()
    norm_headers = [normalize(str(h or "")) for h in headers]
    # 1) correspondances exactes, 2) correspondances par préfixe
    for exact in (True, False):
        for i, h in enumerate(norm_headers):
            if i in mapping or not h:
                continue
            for field, syns in columns.items():
                if field in used:
                    continue
                if any((h == normalize(s)) if exact else (h.startswith(normalize(s) + " ") or normalize(s).startswith(h + " ")) for s in syns):
                    mapping[i] = field
                    used.add(field)
                    break
    return mapping


def rows_from_table(headers: list, rows: list[list], target: str) -> tuple[list[dict], list[str]]:
    mapping = map_headers(headers, target)
    warnings = []
    unmapped = [str(h) for i, h in enumerate(headers) if i not in mapping and h not in (None, "")]
    if unmapped:
        warnings.append("Colonnes non reconnues (ignorées) : " + ", ".join(unmapped))
    if not mapping:
        warnings.append("Aucune colonne reconnue : vérifiez la ligne d'en-tête.")
    out = []
    for n, row in enumerate(rows, start=2):
        if not any(c not in (None, "") for c in row):
            continue
        data = {}
        raw = {}
        for i, field in mapping.items():
            if i < len(row) and row[i] not in (None, ""):
                raw[field] = str(row[i])
                val = convert_value(field, row[i])
                if val not in (None, [], ""):
                    data[field] = val
        out.append({"row_number": n, "data": data, "raw": raw, "evidence": {}, "warnings": []})
    return out, warnings


def read_xlsx(content: bytes) -> tuple[list, list[list]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    # En-tête = première ligne contenant au moins 2 cellules texte
    for idx, r in enumerate(all_rows[:20]):
        if sum(1 for c in r if isinstance(c, str) and c.strip()) >= 2:
            return r, all_rows[idx + 1 :][:5000]
    return [], []


def read_csv(content: bytes) -> tuple[list, list[list]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return [], []
    return rows[0], rows[1:5001]


def text_from_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("PDF protégé par mot de passe.")
    return "\n".join((p.extract_text() or "") for p in reader.pages[:40])


def text_from_docx(content: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(content))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Extraction depuis un texte libre (fiche commerciale, email de demande…)
# ---------------------------------------------------------------------------

NUM = r"(\d{1,3}(?:[  \xa0.]\d{3})+(?:,\d+)?|\d+(?:[.,]\d+)?)"


def _find(pattern: str, text: str, flags=re.IGNORECASE):
    m = re.search(pattern, text, flags)
    if not m:
        return None, None
    start, end = max(0, m.start() - 30), min(len(text), m.end() + 30)
    return m, text[start:end].replace("\n", " ").strip()


def extract_property_from_text(text: str) -> dict:
    data: dict = {}
    ev: dict = {}
    warnings: list[str] = []

    def put(field, value, snippet):
        if value not in (None, "", []):
            data[field] = value
            ev[field] = snippet

    m, s = _find(r"r[ée]f(?:[ée]rence)?\.?\s*(?:annonce|mandat)?\s*[:#n°]*\s*([A-Z0-9][A-Z0-9\-_/]{2,20})", text)
    if m:
        put("reference", m.group(1), s)
    m, s = _find(rf"surface\s*(?:totale|globale|utile|de|d'environ|:|\s)*\s*{NUM}\s*m(?:²|2)", text)
    if m:
        put("surface_total", parse_number(m.group(1)), s)
    else:
        m, s = _find(rf"{NUM}\s*m(?:²|2)", text)
        if m:
            put("surface_total", parse_number(m.group(1)), s)
            warnings.append("Surface totale déduite de la première surface trouvée : à vérifier.")
    m, s = _find(rf"divisible[s]?\s*(?:à partir de|dès|des|a partir de)\s*{NUM}\s*m", text)
    if m:
        put("surface_divisible_min", parse_number(m.group(1)), s)
    m, s = _find(rf"{NUM}\s*m(?:²|2)\s*(?:de\s*)?bureaux", text)
    if m:
        put("surface_offices", parse_number(m.group(1)), s)
    m, s = _find(rf"hauteur\s*(?:libre|sous poutre|utile)?\s*(?:de|:)?\s*{NUM}\s*m(?![²2])", text)
    if m:
        put("clear_height", parse_number(m.group(1)), s)
    m, s = _find(r"(\d{1,3})\s*quais?\b", text)
    if m:
        put("dock_count", int(m.group(1)), s)
    m, s = _find(r"(\d{1,3})\s*portes?\s*sectionnelles?", text)
    if m:
        put("sectional_doors", int(m.group(1)), s)
    m, s = _find(rf"{NUM}\s*kva", text)
    if m:
        put("power_kva", parse_number(m.group(1)), s)
    m, s = _find(rf"{NUM}\s*(?:t|tonnes)\s*/\s*m", text)
    if m:
        put("floor_load", parse_number(m.group(1)), s)
    m, s = _find(r"(\d{1,4})\s*places?\s*de\s*(?:parking|stationnement)", text)
    if m:
        put("parking_spaces", int(m.group(1)), s)
    m, s = _find(r"(acc[eè]s\s*(?:poids[- ]lourds|PL|semi[- ]remorques?))", text)
    if m:
        neg, _ = _find(r"(pas d'|sans |aucun )acc[eè]s\s*(?:poids[- ]lourds|PL)", text)
        put("hgv_access", False if neg else True, s)
    m, s = _find(rf"loyer[^\d\n]{{0,40}}{NUM}\s*(?:€|eur(?:os)?)\s*(?:HT\s*)?(?:HC\s*)?/\s*m", text)
    if m:
        put("rent_per_sqm", parse_number(m.group(1)), s)
    m, s = _find(rf"loyer[^\d\n]{{0,40}}{NUM}\s*(?:€|eur(?:os)?)(?!\s*(?:HT\s*)?(?:HC\s*)?/\s*m)", text)
    if m:
        put("annual_rent", parse_number(m.group(1)), s)
        if not re.search(r"/\s*an|par an|annuel", s, re.IGNORECASE):
            warnings.append("Périodicité du loyer non précisée dans le document : vérifier qu'il s'agit d'un loyer annuel.")
    m, s = _find(rf"prix[^\d\n]{{0,40}}{NUM}\s*(?:€|eur(?:os)?)", text)
    if m:
        put("sale_price", parse_number(m.group(1)), s)
    m, s = _find(rf"taxe\s*fonci[eè]re[^\d\n]{{0,20}}{NUM}\s*(?:€|eur)", text)
    if m:
        put("property_tax", parse_number(m.group(1)), s)
    m, s = _find(rf"charges[^\d\n]{{0,30}}{NUM}\s*(?:€|eur)", text)
    if m:
        put("charges", parse_number(m.group(1)), s)
    m, s = _find(r"honoraires\s*[:\-]?\s*([^\n]{3,80})", text)
    if m:
        put("fees", m.group(1).strip(), s)
    m, s = _find(r"\b(\d{5})\s+([A-ZÉÈÀÂÎÔ][A-Za-zÉÈÊÀÂÎÔÛÇéèêàâîôûç'\- ]{1,40})", text, 0)
    if m:
        put("postal_code", m.group(1), s)
        put("city", m.group(2).strip().split("\n")[0].strip(" -"), s)
    m, s = _find(r"disponib(?:le|ilit[ée])\s*(?:le|au|à partir du|:)?\s*(\d{1,2}/\d{1,2}/\d{2,4})", text)
    if m:
        put("available_from", parse_date(m.group(1)), s)
    elif re.search(r"disponib\w*\s*(imm[ée]diate|de suite|imm[ée]diatement)", text, re.IGNORECASE):
        _, s = _find(r"disponib\w*\s*(imm[ée]diate|de suite|imm[ée]diatement)", text)
        put("availability_note", "Disponibilité immédiate (selon le document)", s)

    head = text[:1500]
    types = parse_types(head)
    if types:
        put("property_type", types[0], head[:120].replace("\n", " "))
        if len(types) > 1:
            warnings.append("Plusieurs types de biens mentionnés : " + ", ".join(types) + ". Vérifiez la typologie.")
    t = parse_transaction(head)
    if t:
        put("transaction", t, head[:120].replace("\n", " "))
    emails = sorted(set(re.findall(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)))
    if emails:
        put("contact_email", emails[0], emails[0])
        if len(emails) > 1:
            warnings.append("Plusieurs emails trouvés : " + ", ".join(emails[:5]))
    phones = re.findall(r"(?:\+33\s?|0)[1-9](?:[\s.-]?\d{2}){4}", text)
    if phones:
        put("contact_phone", phones[0], phones[0])
    desc = re.sub(r"\n{3,}", "\n\n", text).strip()
    if desc:
        data["description"] = desc[:4000]
    return {"data": data, "evidence": ev, "warnings": warnings}


def extract_request_from_text(text: str) -> dict:
    data: dict = {}
    ev: dict = {}
    warnings: list[str] = []

    def put(field, value, snippet):
        if value not in (None, "", []):
            data[field] = value
            ev[field] = snippet

    m, s = _find(rf"(?:entre|de)\s*{NUM}\s*(?:m(?:²|2))?\s*(?:et|à|-)\s*{NUM}\s*m(?:²|2)", text)
    if m:
        put("surface_min", parse_number(m.group(1)), s)
        put("surface_max", parse_number(m.group(2)), s)
    else:
        m, s = _find(rf"(?:minimum|min\.?|au moins|plus de)\s*{NUM}\s*m(?:²|2)", text)
        if m:
            put("surface_min", parse_number(m.group(1)), s)
        m, s = _find(rf"(?:maximum|max\.?|au plus|moins de)\s*{NUM}\s*m(?:²|2)", text)
        if m:
            put("surface_max", parse_number(m.group(1)), s)
    m, s = _find(rf"(?:budget|loyer max\w*|prix max\w*)[^\d\n]{{0,30}}{NUM}\s*(k€|m€|€|eur\w*)", text)
    if m:
        put("budget_max", parse_number(m.group(1) + " " + m.group(2)), s)
    m, s = _find(rf"hauteur\s*(?:libre|sous poutre)?\s*(?:minimum|min\.?|de)?\s*{NUM}\s*m(?![²2])", text)
    if m:
        put("min_clear_height", parse_number(m.group(1)), s)
    m, s = _find(r"(\d{1,3})\s*quais?\b", text)
    if m:
        put("min_docks", int(m.group(1)), s)
    m, s = _find(r"acc[eè]s\s*(?:poids[- ]lourds|PL|semi)", text)
    if m:
        put("requires_hgv_access", True, s)
    m, s = _find(rf"{NUM}\s*kva", text)
    if m:
        put("min_power_kva", parse_number(m.group(1)), s)
    m, s = _find(r"rayon\s*(?:de)?\s*(\d{1,3})\s*km", text)
    if m:
        put("radius_km", float(m.group(1)), s)
    m, s = _find(r"autour de\s+([A-ZÉÈ][\w\-' ]{2,40})", text, 0)
    if m:
        put("center_label", m.group(1).strip(), s)
    depts = re.findall(r"\(?\b(?:d[ée]partement|dpt|dép\.?)\s*(?:du|de la|des)?\s*(\d{2}|2A|2B)\b", text, re.IGNORECASE)
    if depts:
        put("departments", sorted(set(d.upper() for d in depts)), ", ".join(depts))
    cps = sorted(set(re.findall(r"\b(\d{5})\b", text)))
    if cps:
        put("postal_codes", cps[:20], ", ".join(cps[:20]))
    types = parse_types(text[:3000])
    if types:
        put("property_types", types, "types mentionnés")
    t = parse_transaction(text[:3000])
    if t:
        put("transaction", t, "transaction mentionnée")
    m, s = _find(r"(?:d[ée]but|courant|avant|pour|[àa] partir de)\s*(?:le\s*)?(\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2}/\d{4})", text)
    if m:
        put("desired_date", parse_date(m.group(1)), s)
    m, s = _find(r"imp[ée]ratif[s]?\s*:?\s*([^\n]{3,200})", text)
    if m:
        put("imperatives", parse_imperatives(m.group(1)), s)
    emails = sorted(set(re.findall(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)))
    if emails:
        put("email", emails[0], emails[0])
    phones = re.findall(r"(?:\+33\s?|0)[1-9](?:[\s.-]?\d{2}){4}", text)
    if phones:
        put("phone", phones[0], phones[0])
    data["comments"] = text.strip()[:4000]
    warnings.append("Communes, client et interlocuteur ne sont pas déduits automatiquement d'un texte libre : complétez-les.")
    return {"data": data, "evidence": ev, "warnings": warnings}


def extract(file_kind: str, content: bytes, target: str) -> tuple[list[dict], list[str]]:
    """Point d'entrée : retourne (lignes proposées, avertissements globaux)."""
    if file_kind in ("xlsx", "csv"):
        headers, rows = read_xlsx(content) if file_kind == "xlsx" else read_csv(content)
        if not headers:
            return [], ["Aucune ligne d'en-tête détectée."]
        return rows_from_table(headers, rows, target)
    if file_kind == "pdf":
        text = text_from_pdf(content)
    elif file_kind == "docx":
        text = text_from_docx(content)
    elif file_kind == "txt":
        text = content.decode("utf-8", errors="replace")
    else:
        return [], [f"Format {file_kind} non pris en charge pour l'extraction."]
    if len(text.strip()) < 20:
        return [], ["Aucun texte exploitable (document scanné ? L'OCR n'est pas activé) : saisissez les informations manuellement."]
    res = extract_property_from_text(text) if target == "property" else extract_request_from_text(text)
    return [{"row_number": 1, "data": res["data"], "raw": {}, "evidence": res["evidence"], "warnings": res["warnings"]}], []
